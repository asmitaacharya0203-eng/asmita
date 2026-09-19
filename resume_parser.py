"""
parser.py
---------
Reads resumes (PDF / TXT / DOCX) and turns each one into a structured record:
name, e-mail, phone, skills, certifications, education level, years of experience.
"""

import io
import re
from datetime import date
from pathlib import Path

import skills_db

SUPPORTED_EXTENSIONS = (".pdf", ".txt", ".docx")


# --------------------------------------------------------------------------- #
# 1. TEXT EXTRACTION
# --------------------------------------------------------------------------- #
def _read_pdf(data: bytes) -> str:
    """Extract text from a PDF. pdfplumber first, pypdf as a fallback."""
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    except Exception:
        text = ""
    if len(text.strip()) < 30:            # pdfplumber failed / returned nothing -> try pypdf
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            pass
    return text


def _read_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("python-docx is not installed (pip install python-docx) - cannot read .docx files")
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, filename: str) -> str:
    """Return the plain text of a resume given its raw bytes and file name."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".txt":
        text = data.decode("utf-8", errors="ignore")
    elif ext == ".docx":
        text = _read_docx(data)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    # normalise whitespace / odd PDF characters
    text = text.replace("\u00a0", " ").replace("\u2022", "-").replace("\uf0b7", "-")
    return re.sub(r"[ \t]+", " ", text).strip()


# --------------------------------------------------------------------------- #
# 2. SECTION SPLITTING
# --------------------------------------------------------------------------- #
_SECTION_HEADINGS = {
    "experience": r"(?:work|professional|relevant|employment)?\s*(?:experience|employment history|work history|career history|internships?)",
    "education": r"(?:education(?:al)?(?: qualifications?| background)?|academics?|academic (?:qualifications?|background)|qualifications?)",
    "skills": r"(?:(?:technical|key|core|professional|functional)\s+)?(?:skills?|competenc(?:y|ies)|skill set|expertise|tools(?: & | and )?technologies)",
    "certifications": r"(?:certifications?|licen[sc]es?(?: (?:&|and) certifications?)?|professional (?:certifications?|qualifications?)|courses|achievements|awards)",
    "projects": r"(?:projects?|academic projects?|key projects?)",
    "summary": r"(?:(?:professional |career )?summary|profile|objective|career objective|about me)",
}


def split_sections(text: str) -> dict:
    """
    Split resume text into sections by looking for short heading-like lines.
    Returns {section_name: text}. Text before the first heading is stored as 'header'.
    """
    sections = {"header": []}
    current = "header"
    for line in text.splitlines():
        stripped = line.strip().strip(":-_= ").strip()
        matched = None
        if 0 < len(stripped) <= 45:
            for name, pattern in _SECTION_HEADINGS.items():
                if re.fullmatch(pattern, stripped, flags=re.IGNORECASE):
                    matched = name
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


# --------------------------------------------------------------------------- #
# 3. CONTACT DETAILS
# --------------------------------------------------------------------------- #
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s\-()]{8,16}\d)(?!\d)")
_NAME_STOPWORDS = {"resume", "curriculum", "vitae", "cv", "profile", "summary", "objective", "contact", "address",
                   "email", "phone", "mobile", "linkedin", "page", "finance", "analyst", "manager", "associate",
                   "curriculum vitae"}


def extract_email(text: str) -> str:
    m = _EMAIL_RE.search(text)
    return m.group(0) if m else ""


def extract_phone(text: str) -> str:
    for m in _PHONE_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group(1))
        if 10 <= len(digits) <= 13:
            return m.group(1).strip()
    return ""


def extract_name(text: str, filename: str) -> str:
    """Heuristic: the first short, purely alphabetic line near the top of the resume."""
    for line in [l.strip() for l in text.splitlines() if l.strip()][:6]:
        if any(ch.isdigit() for ch in line) or "@" in line:
            continue
        words = line.replace(",", " ").split()
        if not 2 <= len(words) <= 4:
            continue
        if not all(re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*", w) for w in words):
            continue
        if any(w.lower() in _NAME_STOPWORDS for w in words):
            continue
        return " ".join(w.capitalize() if w.isupper() else w for w in words)
    # fall back to the file name:  "rahul_sharma_resume.pdf" -> "Rahul Sharma"
    stem = re.sub(r"(?i)[_\-\s]*(resume|cv)[_\-\s]*", " ", Path(filename).stem)
    return re.sub(r"[_\-]+", " ", stem).strip().title() or Path(filename).stem


# --------------------------------------------------------------------------- #
# 4. YEARS OF EXPERIENCE
# --------------------------------------------------------------------------- #
_MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
           "sep": 9, "oct": 10, "nov": 11, "dec": 12}
_MONTH_RE = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE_RANGE_RE = re.compile(
    rf"(?:(?P<sm>{_MONTH_RE})\.?,?\s+|(?P<sn>\d{{1,2}})[/.]\s*)?(?P<sy>(?:19|20)\d{{2}})"
    rf"\s*(?:-|\u2013|\u2014|to|until|till)\s*"
    rf"(?:(?:(?P<em>{_MONTH_RE})\.?,?\s+|(?P<en>\d{{1,2}})[/.]\s*)?(?P<ey>(?:19|20)\d{{2}})"
    rf"|(?P<present>present|current|currently|till date|till now|now|ongoing|date))",
    re.IGNORECASE,
)
_EXPLICIT_EXP_RE = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+(?:relevant|total|overall|work|professional|industry|post[- ]qualification))*"
    r"\s+(?:experience|exp)\b",
    re.IGNORECASE,
)


def _month_index(year, month):
    return year * 12 + (month - 1)


def _to_month(name, num):
    if name:
        return _MONTHS[name[:3].lower()]
    if num and 1 <= int(num) <= 12:
        return int(num)
    return None


def years_from_date_ranges(text: str, today: date = None) -> float:
    """Sum the (merged, non-overlapping) date ranges such as 'Jun 2019 - Mar 2022' / '2018 - Present'."""
    today = today or date.today()
    now_idx = _month_index(today.year, today.month)
    intervals = []
    for m in _DATE_RANGE_RE.finditer(text):
        sy = int(m.group("sy"))
        sm = _to_month(m.group("sm"), m.group("sn")) or 1
        start = _month_index(sy, sm)
        if m.group("present"):
            end = now_idx
        else:
            ey = int(m.group("ey"))
            em = _to_month(m.group("em"), m.group("en")) or sm
            end = _month_index(ey, em)
        end = min(end, now_idx)            # future dates (e.g. 'expected 2027') are not experience
        if end > start and (end - start) <= 12 * 45:
            intervals.append((start, end))
    if not intervals:
        return 0.0
    intervals.sort()
    merged = [list(intervals[0])]
    for s, e in intervals[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return round(sum(e - s for s, e in merged) / 12, 1)


def extract_experience_years(text: str, sections: dict) -> float:
    """
    Years of experience = the larger of
      (a) an explicit statement such as '5 years of experience', and
      (b) the total duration of the date ranges found in the experience section
          (or, when there is no experience heading, the whole resume minus education).
    """
    explicit = [float(x) for x in _EXPLICIT_EXP_RE.findall(text) if float(x) <= 40]
    explicit_years = max(explicit) if explicit else 0.0

    if sections.get("experience"):
        scope = sections["experience"]
    else:
        scope = text.replace(sections.get("education", ""), "") if sections.get("education") else text
    return max(explicit_years, years_from_date_ranges(scope))


# --------------------------------------------------------------------------- #
# 5. PUTTING IT TOGETHER
# --------------------------------------------------------------------------- #
def parse_resume(text: str, filename: str) -> dict:
    """Convert raw resume text into a structured candidate record."""
    sections = split_sections(text)
    education_text = sections.get("education", "") or text
    certs, certs_in_progress = skills_db.extract_certifications(text)
    return {
        "file": filename,
        "name": extract_name(text, filename),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": skills_db.extract_skills(text),                 # {skill: mentions}
        "certifications": certs,                                  # completed qualifications
        "certs_in_progress": certs_in_progress,                   # e.g. "CFA Level II candidate"
        "education_level": skills_db.extract_education_level(education_text),
        "finance_degree": skills_db.has_finance_degree(education_text),
        "experience_years": extract_experience_years(text, sections),
        "text": text,
    }


def load_resume_file(path) -> dict:
    """Read and parse a single resume from disk."""
    path = Path(path)
    text = extract_text(path.read_bytes(), path.name)
    if len(text) < 30:
        raise ValueError("no extractable text (empty, corrupted or scanned-image file - OCR would be needed)")
    return parse_resume(text, path.name)


def load_resumes_from_folder(folder):
    """
    Read every supported resume in `folder` (sub-folders included).
    Returns (list_of_candidate_records, list_of_(filename, error_message)).
    """
    folder = Path(folder).expanduser()
    if not folder.is_dir():
        raise FileNotFoundError(f"Folder not found: {folder}")
    candidates, errors = [], []
    for path in sorted(folder.rglob("*")):
        if path.suffix.lower() in SUPPORTED_EXTENSIONS and path.is_file():
            try:
                candidates.append(load_resume_file(path))
            except Exception as exc:          # one bad file must not stop the whole batch
                errors.append((path.name, str(exc)))
    return candidates, errors


def load_resumes_from_uploads(uploaded_files):
    """Same as load_resumes_from_folder but for Streamlit UploadedFile objects."""
    candidates, errors = [], []
    for f in uploaded_files:
        try:
            text = extract_text(f.getvalue(), f.name)
            if len(text) < 30:
                raise ValueError("no extractable text (empty, corrupted or scanned-image file - OCR would be needed)")
            candidates.append(parse_resume(text, f.name))
        except Exception as exc:
            errors.append((f.name, str(exc)))
    return candidates, errors
